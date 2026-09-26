"""
Descarga semanal del portal público SíseVe (MINEDU).

Baja el Excel público de casos (el botón «Descargar Excel» del tablero
https://siseve.minedu.gob.pe/Web/App/Mapa) y lo AGREGA en el acto. El Excel
crudo queda en data/raw/ (ignorado por git) y nunca se publica.

Salidas (agregadas, sin datos personales):
  data/estado.json   estado de atención por UGEL (U) y curva de cierre por antigüedad (MAT)
  data/mensual.json  reportes por mes: nacional, por tipo y por UGEL

El Excel público llega hasta UGEL; no trae el colegio. Las cifras por colegio
vienen de la base por transparencia (data/base.json).

    python scripts/siseve_portal.py              descarga y agrega
    python scripts/siseve_portal.py --local X    agrega un Excel ya descargado
"""
import argparse
import datetime as dt
import json
import pathlib
import sys
from collections import Counter, defaultdict

import openpyxl
import requests

ROOT = pathlib.Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
BASE = "https://siseve.minedu.gob.pe/Web"
UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/131.0 Safari/537.36"

ESTADOS = {
    "atención finalizada": "fin",
    "atención finalizada por validar": "proc",
    "atención en proceso": "proc",
    "pendiente de atención por la ie": "pend",
    "observado por ugel": "pend",
}
TIPOS = {"física": "fisica", "psicológica": "psicologica", "sexual": "sexual"}


def descargar_excel(destino: pathlib.Path) -> pathlib.Path:
    s = requests.Session()
    s.headers["User-Agent"] = UA
    s.get(f"{BASE}/App/Mapa", timeout=90).raise_for_status()   # sesión, como el navegador
    r = s.post(f"{BASE}/Inicio/DescargarEXCEL", data="",
               headers={"X-Requested-With": "XMLHttpRequest"}, timeout=600)
    r.raise_for_status()
    if not r.content.startswith(b"PK"):
        raise RuntimeError(f"La respuesta no es un .xlsx ({len(r.content)} bytes): el portal cambió.")
    destino.parent.mkdir(parents=True, exist_ok=True)
    destino.write_bytes(r.content)
    return destino


def _fecha(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    t = str(v).strip()
    for f, n in (("%d/%m/%Y", 10), ("%Y-%m-%d", 10)):
        try:
            return dt.datetime.strptime(t[:n], f).date()
        except ValueError:
            pass
    return None


def leer_casos(xlsx: pathlib.Path):
    """Lee solo 8 columnas: fecha, DRE, UGEL, nivel, tipo de reporte, tipo, subtipo, estado.
    Cualquier otra columna se ignora."""
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    ws = wb["BaseCompleta"] if "BaseCompleta" in wb.sheetnames else wb.worksheets[0]
    for fila in ws.iter_rows(min_row=7, values_only=True):
        if not fila or not fila[0]:
            continue
        f = _fecha(fila[0])
        if not f:
            continue
        yield {"f": f, "ugel": str(fila[2] or "").strip(),
               "tipo": TIPOS.get(str(fila[5] or "").strip().lower()),
               "estado": ESTADOS.get(str(fila[7] or "").strip().lower()),
               "estado_txt": str(fila[7] or "").strip()}


def agregar(casos):
    casos = list(casos)
    if not casos:
        raise RuntimeError("El Excel no trae casos.")
    corte = max(c["f"] for c in casos)
    edad = lambda f: (corte.year - f.year) * 12 + corte.month - f.month  # noqa: E731
    anios = sorted({c["f"].year for c in casos})[-3:]

    u = defaultdict(lambda: {"tot": 0, "y": Counter(), "fin": 0, "proc": 0, "pend": 0})
    mat = defaultdict(lambda: [0, 0])
    mes_nac, mes_tipo, mes_ugel = Counter(), defaultdict(Counter), defaultdict(Counter)
    raros = Counter()
    for c in casos:
        ym = f"{c['f'].year}-{c['f'].month:02d}"
        mes_nac[ym] += 1
        mes_ugel[c["ugel"]][ym] += 1
        if c["tipo"]:
            mes_tipo[c["tipo"]][ym] += 1
        x = u[c["ugel"]]
        x["tot"] += 1
        x["y"][c["f"].year] += 1
        if c["estado"]:
            x[c["estado"]] += 1
        else:
            raros[c["estado_txt"]] += 1
        a = edad(c["f"])
        mat[a][0] += 1
        mat[a][1] += c["estado"] == "fin"

    U = sorted(([k, v["tot"], *[v["y"][a] for a in anios], v["fin"], v["proc"], v["pend"]]
                for k, v in u.items()), key=lambda r: -r[1])
    MAT = [[a, n, round(f / n * 100, 1) if n else 0] for a, (n, f) in sorted(mat.items())]
    return corte, anios, raros, U, MAT, mes_nac, mes_tipo, mes_ugel


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--local", type=pathlib.Path, help="Excel ya descargado")
    a = ap.parse_args()
    xlsx = a.local or descargar_excel(RAW / f"siseve_{dt.date.today():%Y%m%d}.xlsx")
    corte, anios, raros, U, MAT, mn, mt, mu = agregar(leer_casos(xlsx))
    if raros:
        print("Estados no reconocidos:", dict(raros))
    hoy = dt.date.today().isoformat()
    (ROOT / "data" / "estado.json").write_text(json.dumps({
        "corte": corte.isoformat(), "descargado": hoy, "anios": anios,
        "fuente": "Excel público del portal SíseVe (Descargar Excel)",
        "estados_no_reconocidos": sum(raros.values()), "U": U, "MAT": MAT,
    }, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    (ROOT / "data" / "mensual.json").write_text(json.dumps({
        "corte": corte.isoformat(), "descargado": hoy,
        "nac": dict(sorted(mn.items())),
        "tipo": {k: dict(sorted(v.items())) for k, v in mt.items()},
        "ugel": {k: dict(sorted(v.items())) for k, v in sorted(mu.items())},
    }, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"Corte {corte} · {sum(r[1] for r in U):,} casos · {len(U)} UGEL · {len(MAT)} cohortes · años {anios}")


if __name__ == "__main__":
    sys.exit(main())
